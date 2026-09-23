import csv
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

from scripts.ETF池筛选 import 指数聚类筛选


def write_initial_selection(
    path: Path,
    selection_date: date,
    etf_code: str,
    index_code: str,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["日期", "代码", "对标指数", "对标指数代码"],
        )
        writer.writeheader()
        writer.writerow(
            {
                "日期": selection_date.isoformat(),
                "代码": etf_code,
                "对标指数": f"指数{index_code}",
                "对标指数代码": index_code,
            }
        )


def build_selection(
    root: Path,
    rows: tuple[dict[str, str], ...],
    selection_date: date = date(2021, 3, 31),
) -> 指数聚类筛选.SelectionInput:
    return 指数聚类筛选.SelectionInput(
        selection_date=selection_date,
        initial_file=root / "initial.csv",
        index_return_file=root / "returns.csv",
        output_file=root / selection_date.strftime("%Y_%m_%d.xlsx"),
        etf_rows=rows,
    )


class IndexClusteringTests(unittest.TestCase):
    def test_discovery_is_limited_to_requested_frequency_directories(
        self,
    ) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            daily_initial = root / "daily" / "initial"
            monthly_initial = root / "monthly" / "initial"
            daily_returns = root / "daily" / "index_returns" / "window_60"
            monthly_returns = (
                root / "monthly" / "index_returns" / "window_60"
            )
            daily_reports = root / "daily" / "reports"
            monthly_reports = root / "monthly" / "reports"
            file_name = "2021_01_29.csv"
            selection_date = date(2021, 1, 29)

            write_initial_selection(
                daily_initial / file_name,
                selection_date,
                "510001.SH",
                "000001.SH",
            )
            write_initial_selection(
                monthly_initial / file_name,
                selection_date,
                "510002.SH",
                "000002.SH",
            )
            daily_returns.mkdir(parents=True)
            monthly_returns.mkdir(parents=True)
            (daily_returns / file_name).write_text("daily", encoding="utf-8")
            (monthly_returns / file_name).write_text(
                "monthly",
                encoding="utf-8",
            )

            _, daily_selections = (
                指数聚类筛选.discover_selection_inputs(
                    daily_initial,
                    daily_returns,
                    daily_reports,
                )
            )
            _, monthly_selections = (
                指数聚类筛选.discover_selection_inputs(
                    monthly_initial,
                    monthly_returns,
                    monthly_reports,
                )
            )

            self.assertEqual(len(daily_selections), 1)
            self.assertEqual(len(monthly_selections), 1)
            self.assertEqual(daily_selections[0].index_codes, {"000001.SH"})
            self.assertEqual(
                monthly_selections[0].index_codes,
                {"000002.SH"},
            )
            self.assertEqual(
                daily_selections[0].index_return_file,
                daily_returns / file_name,
            )
            self.assertEqual(
                monthly_selections[0].index_return_file,
                monthly_returns / file_name,
            )
            self.assertEqual(
                daily_selections[0].output_file,
                daily_reports / "2021_01_29.xlsx",
            )
            self.assertEqual(
                monthly_selections[0].output_file,
                monthly_reports / "2021_01_29.xlsx",
            )

    def test_turnover_window_ends_on_selection_date_without_future_data(
        self,
    ) -> None:
        selection_date = date(2021, 1, 25)
        available_dates = [
            date(2021, 1, 1) + timedelta(days=offset)
            for offset in range(25)
        ]
        selection = build_selection(
            Path("unused"),
            (
                {
                    "代码": "510001.SH",
                    "对标指数代码": "000001.SH",
                },
            ),
            selection_date,
        )

        windows = 指数聚类筛选.build_turnover_windows(
            [selection],
            available_dates,
        )
        window = windows[selection.file_name]

        self.assertEqual(len(window), 20)
        self.assertEqual(window[-1], selection_date)
        self.assertTrue(all(day <= selection_date for day in window))

    def test_correlation_threshold_changes_complete_linkage_result(
        self,
    ) -> None:
        index_codes = ["A", "B", "C"]
        distance_matrix = np.asarray(
            [
                [0.0, 0.2, 0.9],
                [0.2, 0.0, 0.9],
                [0.9, 0.9, 0.0],
            ]
        )

        threshold_07 = 指数聚类筛选.complete_linkage_clusters(
            index_codes,
            distance_matrix,
            0.7,
        )
        threshold_09 = 指数聚类筛选.complete_linkage_clusters(
            index_codes,
            distance_matrix,
            0.9,
        )

        self.assertEqual(threshold_07["A"], threshold_07["B"])
        self.assertNotEqual(threshold_07["A"], threshold_07["C"])
        self.assertNotEqual(threshold_09["A"], threshold_09["B"])
        self.assertEqual(len(set(threshold_09.values())), 3)

    def test_representative_tie_break_and_detail_thresholds(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            rows = (
                {
                    "代码": "510002.SH",
                    "名称": "ETF二",
                    "类别大类": "股票型ETF",
                    "对标指数代码": "A",
                    "对标指数": "指数A",
                },
                {
                    "代码": "510001.SH",
                    "名称": "ETF一",
                    "类别大类": "股票型ETF",
                    "对标指数代码": "B",
                    "对标指数": "指数B",
                },
            )
            selection = build_selection(root, rows)
            cluster_by_code = {"A": 1, "B": 1}
            average_turnover = {
                (selection.file_name, "510001.SH"): 100.0,
                (selection.file_name, "510002.SH"): 100.0,
            }

            selected_rows = 指数聚类筛选.select_cluster_representatives(
                selection,
                cluster_by_code,
                average_turnover,
            )
            detail_rows = 指数聚类筛选.build_cluster_detail_rows(
                selection,
                cluster_by_code,
                average_turnover,
                selected_rows,
                ["A", "B"],
                np.asarray([[1.0, 0.8], [0.8, 1.0]]),
                0.7,
            )

            self.assertEqual(selected_rows[0]["代码"], "510001.SH")
            self.assertEqual(selected_rows[0]["聚类指数数量"], "2")
            self.assertTrue(
                all(float(row["相关性阈值"]) == 0.7 for row in detail_rows)
            )
            self.assertTrue(
                all(
                    abs(float(row["距离阈值"]) - 0.3) < 1e-12
                    for row in detail_rows
                )
            )

    def test_legacy_selection_date_column_remains_readable(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            selection_date = date(2021, 3, 31)
            selection = build_selection(
                root,
                (
                    {
                        "代码": "510001.SH",
                        "对标指数代码": "000001.SH",
                    },
                ),
                selection_date,
            )
            with selection.index_return_file.open(
                "w",
                encoding="utf-8-sig",
                newline="",
            ) as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=[
                        "月末交易日",
                        "收益日期",
                        "对标指数代码",
                        "日收益率",
                    ],
                )
                writer.writeheader()
                for offset in range(60):
                    writer.writerow(
                        {
                            "月末交易日": selection_date.isoformat(),
                            "收益日期": (
                                selection_date - timedelta(days=59 - offset)
                            ).isoformat(),
                            "对标指数代码": "000001.SH",
                            "日收益率": "0.001",
                        }
                    )

            returns_by_code = (
                指数聚类筛选.read_selection_index_returns(selection)
            )

            self.assertEqual(set(returns_by_code), {"000001.SH"})
            self.assertEqual(len(returns_by_code["000001.SH"]), 60)

    def test_cleanup_only_removes_stale_xlsx_in_target_directory(
        self,
    ) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            target = root / "daily" / "threshold_0.8" / "reports"
            other_frequency = (
                root / "monthly" / "threshold_0.8" / "reports"
            )
            other_threshold = (
                root / "daily" / "threshold_0.9" / "reports"
            )
            for directory in (target, other_frequency, other_threshold):
                directory.mkdir(parents=True)

            keep = target / "2021_01_30.xlsx"
            stale = target / "2021_01_29.xlsx"
            unrelated_workbook = target / "说明.xlsx"
            non_padded_workbook = target / "2021_1_2.xlsx"
            legacy_csv = target / "legacy.csv"
            note = target / "说明.txt"
            other_frequency_file = other_frequency / "2021_01_29.xlsx"
            other_threshold_file = other_threshold / "2021_01_29.xlsx"
            for path in (
                keep,
                stale,
                unrelated_workbook,
                non_padded_workbook,
                legacy_csv,
                note,
                other_frequency_file,
                other_threshold_file,
            ):
                path.write_text("test", encoding="utf-8")

            removed_count = 指数聚类筛选.remove_stale_outputs(
                target,
                {keep.name},
            )

            self.assertEqual(removed_count, 1)
            self.assertTrue(keep.exists())
            self.assertFalse(stale.exists())
            self.assertTrue(unrelated_workbook.exists())
            self.assertTrue(non_padded_workbook.exists())
            self.assertTrue(legacy_csv.exists())
            self.assertTrue(note.exists())
            self.assertTrue(other_frequency_file.exists())
            self.assertTrue(other_threshold_file.exists())

    def test_workbook_keeps_three_existing_sheets_and_core_columns(
        self,
    ) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            row = {
                "日期": "2021-03-31",
                "代码": "510001.SH",
                "名称": "ETF一",
                "类别大类": "股票型ETF",
                "对标指数代码": "A",
                "对标指数": "指数A",
            }
            selection = build_selection(root, (row,))
            selected_row = {
                **row,
                "过去20个交易日平均成交额": "100",
                "聚类编号": "1",
                "聚类指数数量": "1",
            }
            detail_row = {
                "聚类编号": "1",
                "是否最终代表": "是",
                "ETF代码": "510001.SH",
                "ETF名称": "ETF一",
                "类别大类": "股票型ETF",
                "对标指数代码": "A",
                "对标指数": "指数A",
                "过去20个交易日平均成交额": "100",
                "与代表指数相关性": "1",
                "聚类内最低相关性": "",
                "相关性阈值": "0.7",
                "距离阈值": "0.3",
            }

            指数聚类筛选.write_selection_workbook(
                selection.output_file,
                list(row),
                [selected_row],
                selection,
                ["A"],
                np.asarray([[1.0]]),
                [detail_row],
            )

            workbook = load_workbook(selection.output_file, read_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["动态指数池", "相关性矩阵", "聚类明细"],
                )
                selection_headers = [
                    cell.value for cell in workbook["动态指数池"][1]
                ]
                self.assertEqual(
                    selection_headers[-3:],
                    指数聚类筛选.ADDED_OUTPUT_COLUMNS,
                )
                detail_headers = [
                    cell.value for cell in workbook["聚类明细"][1]
                ]
                self.assertEqual(
                    detail_headers,
                    指数聚类筛选.CLUSTER_DETAIL_COLUMNS,
                )
            finally:
                workbook.close()

    def test_main_routes_frequency_and_threshold_to_all_paths(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            initial_dir = root / "monthly" / "initial"
            returns_dir = root / "monthly" / "index_returns" / "window_60"
            reports_dir = (
                root / "monthly" / "clusters" / "threshold_0.7" / "reports"
            )
            selection = 指数聚类筛选.SelectionInput(
                selection_date=date(2021, 1, 29),
                initial_file=initial_dir / "2021_01_29.csv",
                index_return_file=returns_dir / "2021_01_29.csv",
                output_file=reports_dir / "2021_01_29.xlsx",
                etf_rows=(
                    {
                        "日期": "2021-01-29",
                        "代码": "510001.SH",
                        "对标指数代码": "A",
                    },
                ),
            )
            returns_by_code = {"A": {date(2021, 1, 29): 0.01}}
            correlation_matrix = np.asarray([[1.0]])
            distance_matrix = np.asarray([[0.0]])
            selected_rows = [{"代码": "510001.SH"}]
            detail_rows = [{"ETF代码": "510001.SH"}]
            with patch.object(
                指数聚类筛选,
                "etf_pool_initial_dir",
                return_value=initial_dir,
            ) as initial_path, patch.object(
                指数聚类筛选,
                "index_return_dir",
                return_value=returns_dir,
            ) as return_path, patch.object(
                指数聚类筛选,
                "cluster_report_dir",
                return_value=reports_dir,
            ) as report_path, patch.object(
                指数聚类筛选,
                "discover_selection_inputs",
                return_value=(["日期", "代码", "对标指数代码"], [selection]),
            ) as discover, patch.object(
                指数聚类筛选,
                "inspect_etf_data_dates",
                return_value=[],
            ), patch.object(
                指数聚类筛选,
                "build_turnover_windows",
                return_value={selection.file_name: ()},
            ), patch.object(
                指数聚类筛选,
                "calculate_average_turnover",
                return_value={},
            ), patch.object(
                指数聚类筛选,
                "read_selection_index_returns",
                return_value=returns_by_code,
            ), patch.object(
                指数聚类筛选,
                "calculate_correlation_and_distance_matrices",
                return_value=(correlation_matrix, distance_matrix),
            ), patch.object(
                指数聚类筛选,
                "complete_linkage_clusters",
                return_value={"A": 1},
            ) as cluster, patch.object(
                指数聚类筛选,
                "select_cluster_representatives",
                return_value=selected_rows,
            ), patch.object(
                指数聚类筛选,
                "build_cluster_detail_rows",
                return_value=detail_rows,
            ) as build_details, patch.object(
                指数聚类筛选,
                "write_selection_workbook",
            ) as write_workbook, patch.object(
                指数聚类筛选,
                "remove_stale_outputs",
                return_value=0,
            ) as cleanup, patch("builtins.print"):
                指数聚类筛选.main("monthly", 0.7)

            initial_path.assert_called_once_with("monthly")
            return_path.assert_called_once_with("monthly", 60)
            report_path.assert_called_once_with("monthly", 0.7)
            discover.assert_called_once_with(
                initial_dir,
                returns_dir,
                reports_dir,
            )
            cluster_args = cluster.call_args.args
            self.assertEqual(cluster_args[0], ["A"])
            self.assertIs(cluster_args[1], distance_matrix)
            self.assertEqual(cluster_args[2], 0.7)
            detail_args = build_details.call_args.args
            self.assertIs(detail_args[0], selection)
            self.assertEqual(detail_args[-1], 0.7)
            write_args = write_workbook.call_args.args
            self.assertEqual(write_args[0], selection.output_file)
            self.assertIs(write_args[3], selection)
            cleanup.assert_called_once_with(
                reports_dir,
                {selection.file_name},
            )

    def test_invalid_parameters_fail_before_path_resolution(self) -> None:
        invalid_cases = (("weekly", 0.9), ("daily", 0.75))
        for update_frequency, correlation_threshold in invalid_cases:
            with self.subTest(
                update_frequency=update_frequency,
                correlation_threshold=correlation_threshold,
            ), patch.object(
                指数聚类筛选,
                "etf_pool_initial_dir",
                side_effect=AssertionError("不应计算初筛路径"),
            ), patch.object(
                指数聚类筛选,
                "index_return_dir",
                side_effect=AssertionError("不应计算收益率路径"),
            ), patch.object(
                指数聚类筛选,
                "cluster_report_dir",
                side_effect=AssertionError("不应计算聚类输出路径"),
            ):
                with self.assertRaises(ValueError):
                    指数聚类筛选.main(
                        update_frequency,
                        correlation_threshold,
                    )


if __name__ == "__main__":
    unittest.main()
