import csv
from collections import defaultdict
from datetime import date, timedelta
import math
import os
from pathlib import Path
from tempfile import TemporaryDirectory, gettempdir
from types import SimpleNamespace
import unittest
from unittest.mock import call, patch

os.environ.setdefault(
    "MPLCONFIGDIR",
    str(Path(gettempdir()) / "etf_rotation_matplotlib_tests"),
)

from openpyxl import load_workbook

from 实验配置 import (
    CORRELATION_THRESHOLDS,
    FILTER_PROFILE_SETTINGS,
    STAGGER_DAYS,
    TREND_FACTORS,
    TREND_WINDOWS,
    UPDATE_FREQUENCIES,
    ExperimentCase,
)
from scripts.ETF趋势策略回测 import 交易策略回测


def make_case(
    update_frequency: str = "daily",
    filter_profile: str = "base",
    correlation_threshold: float = 0.9,
    trend_factor: str = "return_r2",
    trend_window: int = 20,
    stagger_days: int = 1,
) -> ExperimentCase:
    return ExperimentCase(
        update_frequency=update_frequency,
        filter_profile=filter_profile,
        correlation_threshold=correlation_threshold,
        trend_factor=trend_factor,
        trend_window=trend_window,
        stagger_days=stagger_days,
    )


def make_member(index_code: str) -> 交易策略回测.FactorMember:
    return 交易策略回测.FactorMember(
        signal_date=date(2021, 3, 31),
        index_code=index_code,
        index_name=f"指数{index_code}",
        trend_factor=1.0,
        window_return=0.1,
    )


def write_factor_file(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "日期",
                "对标指数代码",
                "对标指数",
                "窗口收益率",
                "趋势质量因子",
                "风险调整趋势得分",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "日期": "2021-03-31",
                "对标指数代码": "A",
                "对标指数": "指数A",
                "窗口收益率": 0.1,
                "趋势质量因子": 2.0,
                "风险调整趋势得分": 1.0,
            }
        )
        writer.writerow(
            {
                "日期": "2021-03-31",
                "对标指数代码": "B",
                "对标指数": "指数B",
                "窗口收益率": 0.1,
                "趋势质量因子": 1.0,
                "风险调整趋势得分": 2.0,
            }
        )


def make_target(
    signal_date: date,
    etf_code: str,
    index_code: str,
) -> 交易策略回测.DailyTarget:
    member = 交易策略回测.TargetMember(
        signal_date=signal_date,
        index_code=index_code,
        index_name=f"指数{index_code}",
        etf_code=etf_code,
        etf_name=f"ETF{etf_code}",
        target_weight=1.0,
        trend_factor=1.0,
        factor_rank=1,
        selection_volume=1000.0,
        selection_amount=100000.0,
        selection_scale=1000000.0,
    )
    return 交易策略回测.DailyTarget(
        signal_date=signal_date,
        planned_index_count=1,
        selected_index_count=1,
        members=(member,),
    )


class ExplodingPriceHistory(dict):
    def get(self, key: object, default: object = None) -> object:
        raise AssertionError("base状态不应读取过滤行情")


class StrategyBacktestTests(unittest.TestCase):
    def test_all_central_case_values_are_valid_and_route_input_paths(
        self,
    ) -> None:
        for update_frequency in UPDATE_FREQUENCIES:
            for correlation_threshold in CORRELATION_THRESHOLDS:
                for trend_window in TREND_WINDOWS:
                    experiment_case = make_case(
                        update_frequency=update_frequency,
                        correlation_threshold=correlation_threshold,
                        trend_window=trend_window,
                    )
                    factor_directory = Path("factor_directory")
                    index_directory = Path("index_directory")
                    with self.subTest(
                        update_frequency=update_frequency,
                        correlation_threshold=correlation_threshold,
                        trend_window=trend_window,
                    ), patch.object(
                        交易策略回测,
                        "factor_dir",
                        return_value=factor_directory,
                    ) as factor_path, patch.object(
                        交易策略回测,
                        "index_price_dir",
                        return_value=index_directory,
                    ) as index_path:
                        settings = 交易策略回测.build_run_settings(
                            experiment_case
                        )

                    factor_path.assert_called_once_with(
                        update_frequency,
                        correlation_threshold,
                        trend_window,
                    )
                    index_path.assert_called_once_with(
                        update_frequency,
                        correlation_threshold,
                    )
                    self.assertEqual(
                        settings.factor_directory,
                        factor_directory,
                    )
                    self.assertEqual(
                        settings.index_price_directory,
                        index_directory,
                    )

        for filter_profile in FILTER_PROFILE_SETTINGS:
            交易策略回测.validate_experiment_case(
                make_case(filter_profile=filter_profile)
            )
        for trend_factor in TREND_FACTORS:
            交易策略回测.validate_experiment_case(
                make_case(trend_factor=trend_factor)
            )
        for stagger_days in STAGGER_DAYS:
            交易策略回测.validate_experiment_case(
                make_case(stagger_days=stagger_days)
            )

    def test_four_filter_profiles_map_to_expected_switches(self) -> None:
        expected = {
            "base": ("none", False),
            "bias_only": ("upper", False),
            "volatility_only": ("none", True),
            "bias_and_volatility": ("upper", True),
        }

        for profile, switches in expected.items():
            with self.subTest(profile=profile):
                settings = 交易策略回测.build_run_settings(
                    make_case(filter_profile=profile)
                )
                self.assertEqual(
                    (settings.bias_mode, settings.vol_filter_enabled),
                    switches,
                )

    def test_invalid_case_fails_before_path_resolution(self) -> None:
        invalid_cases = (
            make_case(update_frequency="weekly"),
            make_case(filter_profile="unknown"),
            make_case(correlation_threshold=0.75),
            make_case(trend_factor="slope"),
            make_case(trend_window=30),
            make_case(stagger_days=0),
        )

        for experiment_case in invalid_cases:
            with self.subTest(experiment_case=experiment_case), patch.object(
                交易策略回测,
                "factor_dir",
                side_effect=AssertionError("不应计算因子路径"),
            ), patch.object(
                交易策略回测,
                "index_price_dir",
                side_effect=AssertionError("不应计算行情路径"),
            ):
                with self.assertRaises(ValueError):
                    交易策略回测.build_run_settings(experiment_case)

    def test_existing_tuning_calls_keep_optional_defaults(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            index_directory = Path(temporary_directory)
            index_file = index_directory / "2021_03_31.csv"
            with index_file.open(
                "w",
                encoding="utf-8-sig",
                newline="",
            ) as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["收益日期", "对标指数代码", "收盘价"],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "收益日期": "2021-03-31",
                        "对标指数代码": "A",
                        "收盘价": 100.0,
                    }
                )

            with patch.object(
                交易策略回测,
                "INDEX_RAW_DATA_DIR",
                index_directory,
            ):
                history = 交易策略回测.load_filter_price_history({"A"})

        self.assertEqual(history["A"], ([date(2021, 3, 31)], [100.0]))
        with patch.object(
            交易策略回测,
            "BIAS_MODE",
            "upper",
        ), patch.object(
            交易策略回测,
            "BIAS_UPPER",
            0.09,
        ):
            self.assertFalse(交易策略回测.passes_bias_filter(0.10))

    def test_filter_profiles_are_distinct_and_ignore_future_prices(
        self,
    ) -> None:
        signal_date = date(2021, 3, 31)
        price_dates = [
            signal_date - timedelta(days=35 - offset)
            for offset in range(36)
        ]
        members = tuple(make_member(code) for code in ("A", "B", "C"))
        prices = {
            "A": [
                95.0 if offset % 2 == 0 else 105.0
                for offset in range(35)
            ]
            + [150.0],
            "B": [
                90.0 if offset % 2 == 0 else 110.0
                for offset in range(35)
            ]
            + [100.0],
            "C": [100.0 + 0.01 * offset for offset in range(36)],
        }
        price_history = {
            code: (price_dates, values)
            for code, values in prices.items()
        }
        history_with_future = {
            code: (
                [*price_dates, signal_date + timedelta(days=1)],
                [*values, 1000000.0],
            )
            for code, values in prices.items()
        }
        expected_codes = {
            "base": ("A", "B", "C"),
            "bias_only": ("B", "C"),
            "volatility_only": ("A", "B"),
            "bias_and_volatility": ("B",),
        }

        base_settings = 交易策略回测.build_run_settings(
            make_case(filter_profile="base")
        )
        base_members = 交易策略回测.filter_eligible_members(
            signal_date,
            members,
            ExplodingPriceHistory(),
            base_settings,
        )
        self.assertEqual(
            tuple(member.index_code for member in base_members),
            expected_codes["base"],
        )

        for profile in (
            "bias_only",
            "volatility_only",
            "bias_and_volatility",
        ):
            with self.subTest(profile=profile):
                settings = 交易策略回测.build_run_settings(
                    make_case(filter_profile=profile)
                )
                selected = 交易策略回测.filter_eligible_members(
                    signal_date,
                    members,
                    price_history,
                    settings,
                )
                selected_with_future = (
                    交易策略回测.filter_eligible_members(
                        signal_date,
                        members,
                        history_with_future,
                        settings,
                    )
                )
                codes = tuple(member.index_code for member in selected)
                future_codes = tuple(
                    member.index_code for member in selected_with_future
                )
                self.assertEqual(codes, expected_codes[profile])
                self.assertEqual(future_codes, codes)

    def test_factor_columns_produce_independent_rankings(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            factor_directory = root / "factors"
            write_factor_file(factor_directory / "2021.csv")
            base_settings = 交易策略回测.build_run_settings(make_case())
            settings = 交易策略回测.BacktestRunSettings(
                experiment_case=base_settings.experiment_case,
                factor_directory=factor_directory,
                index_price_directory=root / "index_prices",
                bias_mode=base_settings.bias_mode,
                vol_filter_enabled=base_settings.vol_filter_enabled,
            )

            with patch.object(
                交易策略回测,
                "load_filter_price_history",
                side_effect=AssertionError("base状态不应加载过滤行情"),
            ):
                return_r2 = 交易策略回测.read_daily_top_factors(
                    "趋势质量因子",
                    settings,
                )
                return_vol = 交易策略回测.read_daily_top_factors(
                    "风险调整趋势得分",
                    settings,
                )

            signal_date = date(2021, 3, 31)
            self.assertEqual(
                return_r2[signal_date].members[0].index_code,
                "A",
            )
            self.assertEqual(
                return_vol[signal_date].members[0].index_code,
                "B",
            )

    def test_filtered_factor_reader_uses_case_index_price_directory(
        self,
    ) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            factor_directory = root / "factors"
            index_directory = root / "monthly_index_prices"
            write_factor_file(factor_directory / "2021.csv")
            original_settings = 交易策略回测.build_run_settings(
                make_case(
                    update_frequency="monthly",
                    filter_profile="bias_only",
                )
            )
            settings = 交易策略回测.BacktestRunSettings(
                experiment_case=original_settings.experiment_case,
                factor_directory=factor_directory,
                index_price_directory=index_directory,
                bias_mode=original_settings.bias_mode,
                vol_filter_enabled=original_settings.vol_filter_enabled,
            )
            signal_date = date(2021, 3, 31)
            price_dates = [
                signal_date - timedelta(days=35 - offset)
                for offset in range(36)
            ]
            history = {
                code: (price_dates, [100.0] * 36)
                for code in ("A", "B")
            }

            with patch.object(
                交易策略回测,
                "load_filter_price_history",
                return_value=history,
            ) as load_history:
                selections = 交易策略回测.read_daily_top_factors(
                    "趋势质量因子",
                    settings,
                )

            load_history.assert_called_once_with(
                {"A", "B"},
                index_directory,
            )
            self.assertIn(signal_date, selections)

    def test_all_stagger_intervals_control_accounts_and_schedule(self) -> None:
        trading_dates = [
            date(2021, 1, 1) + timedelta(days=offset)
            for offset in range(21)
        ]

        for stagger_days in range(1, 11):
            with self.subTest(stagger_days=stagger_days):
                settings = 交易策略回测.build_run_settings(
                    make_case(stagger_days=stagger_days)
                )
                (
                    _,
                    nav_records,
                    account_daily_records,
                    _,
                    _,
                ) = 交易策略回测.run_backtest(
                    "close",
                    trading_dates,
                    {},
                    {},
                    settings,
                )

                self.assertEqual(len(nav_records), len(trading_dates))
                rows_by_date: dict[date, list[object]] = defaultdict(list)
                attempted_positions: dict[int, list[int]] = defaultdict(list)
                position_by_date = {
                    current_date: position
                    for position, current_date in enumerate(trading_dates)
                }
                for row in account_daily_records:
                    rows_by_date[row.current_date].append(row)
                    if row.rebalance_attempted:
                        attempted_positions[row.account_id].append(
                            position_by_date[row.current_date]
                        )

                self.assertTrue(
                    all(
                        len(rows) == stagger_days
                        for rows in rows_by_date.values()
                    )
                )
                first_day_rows = rows_by_date[trading_dates[0]]
                self.assertTrue(
                    all(
                        math.isclose(
                            row.account_nav,
                            1.0 / stagger_days,
                            abs_tol=1e-12,
                        )
                        for row in first_day_rows
                    )
                )
                self.assertEqual(
                    set(attempted_positions),
                    set(range(1, stagger_days + 1)),
                )
                for account_id, positions in attempted_positions.items():
                    self.assertEqual(positions[0], account_id - 1)
                    self.assertTrue(
                        all(
                            later - earlier == stagger_days
                            for earlier, later in zip(
                                positions,
                                positions[1:],
                            )
                        )
                    )

                next_day_result = 交易策略回测.run_backtest(
                    "next_day_vwap",
                    trading_dates,
                    {},
                    {},
                    settings,
                )
                next_day_attempts: dict[int, list[int]] = defaultdict(list)
                for row in next_day_result[2]:
                    if row.rebalance_attempted:
                        next_day_attempts[row.account_id].append(
                            position_by_date[row.current_date]
                        )
                self.assertEqual(
                    set(next_day_attempts),
                    set(range(1, stagger_days + 1)),
                )
                for account_id, positions in next_day_attempts.items():
                    self.assertEqual(positions[0], account_id)
                    self.assertTrue(
                        all(
                            later - earlier == stagger_days
                            for earlier, later in zip(
                                positions,
                                positions[1:],
                            )
                        )
                    )

    def test_close_and_next_day_vwap_use_correct_signal_dates(self) -> None:
        first_date = date(2021, 1, 4)
        second_date = date(2021, 1, 5)
        third_date = date(2021, 1, 6)
        trading_dates = [first_date, second_date, third_date]
        targets = {
            first_date: make_target(first_date, "ETF1", "INDEX1"),
            second_date: make_target(second_date, "ETF2", "INDEX2"),
        }
        prices = {
            current_date: {
                "ETF1": 交易策略回测.PricePoint(
                    close=100.0,
                    vwap=101.0,
                    amount=1000000.0,
                ),
                "ETF2": 交易策略回测.PricePoint(
                    close=200.0,
                    vwap=202.0,
                    amount=1000000.0,
                ),
            }
            for current_date in trading_dates
        }
        settings = 交易策略回测.build_run_settings(make_case())

        close_result = 交易策略回测.run_backtest(
            "close",
            trading_dates,
            targets,
            prices,
            settings,
        )
        vwap_result = 交易策略回测.run_backtest(
            "next_day_vwap",
            trading_dates,
            targets,
            prices,
            settings,
        )
        close_buys = [
            (trade.execution_date, trade.signal_date, trade.etf_code)
            for trade in close_result[4]
            if trade.direction == "买入"
        ]
        vwap_buys = [
            (trade.execution_date, trade.signal_date, trade.etf_code)
            for trade in vwap_result[4]
            if trade.direction == "买入"
        ]

        self.assertEqual(
            close_buys,
            [
                (first_date, first_date, "ETF1"),
                (second_date, second_date, "ETF2"),
            ],
        )
        self.assertEqual(
            vwap_buys,
            [
                (second_date, first_date, "ETF1"),
                (third_date, second_date, "ETF2"),
            ],
        )
        expected_second_day_nav = (
            1.0
            / (1.0 + 交易策略回测.TRANSACTION_COST_RATE)
            * 100.0
            / 101.0
        )
        actual_second_day_nav = next(
            record.nav
            for record in vwap_result[1]
            if record.current_date == second_date
        )
        self.assertAlmostEqual(
            actual_second_day_nav,
            expected_second_day_nav,
            places=12,
        )

    def test_main_maps_return_r2_to_its_own_factor_column(self) -> None:
        benchmark = 交易策略回测.BenchmarkData(
            code="BENCHMARK",
            name="测试基准",
            closes={date(2021, 3, 31): 100.0},
        )
        with patch.object(
            交易策略回测,
            "load_benchmark_data",
            return_value=benchmark,
        ), patch.object(
            交易策略回测,
            "read_daily_top_factors",
            side_effect=RuntimeError("stop after factor routing"),
        ) as read_factors, patch("builtins.print"):
            with self.assertRaisesRegex(
                RuntimeError,
                "stop after factor routing",
            ):
                交易策略回测.main(
                    "daily",
                    "base",
                    0.7,
                    "return_r2",
                    10,
                    1,
                )

        self.assertEqual(
            read_factors.call_args.args[0],
            "趋势质量因子",
        )

    def test_main_runs_one_factor_and_routes_both_execution_modes(
        self,
    ) -> None:
        signal_date = date(2021, 3, 31)
        experiment_case = make_case(
            update_frequency="monthly",
            filter_profile="bias_and_volatility",
            correlation_threshold=0.8,
            trend_factor="return_vol",
            trend_window=60,
            stagger_days=10,
        )
        benchmark = 交易策略回测.BenchmarkData(
            code="BENCHMARK",
            name="测试基准",
            closes={signal_date: 100.0},
        )
        selection = 交易策略回测.DailyFactorSelection(
            signal_date=signal_date,
            planned_index_count=0,
            members=(),
        )
        target = 交易策略回测.DailyTarget(
            signal_date=signal_date,
            planned_index_count=0,
            selected_index_count=0,
            members=(),
        )
        performance = {
            "年化收益率": 0.1,
            "夏普比率": 1.0,
            "最大回撤": -0.1,
            "平均每日单边换手率": 0.2,
        }
        relative_performance = {
            "超额累计收益率": 0.05,
            "信息比率": 0.5,
        }
        output_root = Path("isolated_backtest_outputs")

        with patch.object(
            交易策略回测,
            "load_benchmark_data",
            return_value=benchmark,
        ), patch.object(
            交易策略回测,
            "read_daily_top_factors",
            return_value={signal_date: selection},
        ) as read_factors, patch.object(
            交易策略回测,
            "build_daily_targets",
            return_value=({signal_date: target}, [signal_date]),
        ), patch.object(
            交易策略回测,
            "load_selected_prices",
            return_value={},
        ), patch.object(
            交易策略回测,
            "backtest_result_dir",
            side_effect=lambda case_value, mode: output_root / mode,
        ) as result_path, patch.object(
            交易策略回测,
            "run_backtest",
            return_value=([], [object()], [], [], []),
        ) as run_backtest, patch.object(
            交易策略回测,
            "build_benchmark_records",
            return_value=(signal_date, [object()]),
        ), patch.object(
            交易策略回测,
            "write_mode_workbooks",
            side_effect=lambda *args: [args[10] / "result.xlsx"],
        ) as write_workbooks, patch.object(
            交易策略回测,
            "write_mode_figures",
            side_effect=lambda *args: [args[4] / "result.png"],
        ) as write_figures, patch.object(
            交易策略回测,
            "validate_mode_outputs",
        ) as validate_outputs, patch.object(
            交易策略回测,
            "calculate_performance",
            return_value=performance,
        ), patch.object(
            交易策略回测,
            "calculate_relative_performance",
            return_value=relative_performance,
        ), patch("builtins.print"):
            交易策略回测.main(
                experiment_case.update_frequency,
                experiment_case.filter_profile,
                experiment_case.correlation_threshold,
                experiment_case.trend_factor,
                experiment_case.trend_window,
                experiment_case.stagger_days,
            )

        self.assertEqual(read_factors.call_count, 1)
        self.assertEqual(
            read_factors.call_args.args[0],
            "风险调整趋势得分",
        )
        routed_settings = read_factors.call_args.args[1]
        self.assertEqual(
            routed_settings.experiment_case,
            experiment_case,
        )
        self.assertEqual(
            (routed_settings.bias_mode, routed_settings.vol_filter_enabled),
            ("upper", True),
        )
        self.assertEqual(
            routed_settings.factor_directory,
            交易策略回测.PROJECT_ROOT
            / "outputs"
            / "etf_trend_strategy"
            / "monthly"
            / "threshold_0.8"
            / "factors"
            / "window_60",
        )
        self.assertEqual(
            routed_settings.index_price_directory,
            交易策略回测.PROJECT_ROOT
            / "outputs"
            / "etf_trend_strategy"
            / "monthly"
            / "threshold_0.8"
            / "index_prices",
        )
        self.assertEqual(
            result_path.call_args_list,
            [
                call(experiment_case, "close"),
                call(experiment_case, "next_day_vwap"),
            ],
        )
        self.assertEqual(
            [item.args[0] for item in run_backtest.call_args_list],
            ["close", "next_day_vwap"],
        )
        self.assertTrue(
            all(
                item.args[4] == routed_settings
                for item in run_backtest.call_args_list
            )
        )
        self.assertEqual(
            [item.args[10] for item in write_workbooks.call_args_list],
            [output_root / "close", output_root / "next_day_vwap"],
        )
        self.assertTrue(
            all(
                item.args[11] == routed_settings
                for item in write_workbooks.call_args_list
            )
        )
        self.assertEqual(
            [item.args[4] for item in write_figures.call_args_list],
            [output_root / "close", output_root / "next_day_vwap"],
        )
        self.assertEqual(validate_outputs.call_count, 2)
        self.assertTrue(
            all(
                item.args[9] == routed_settings
                for item in validate_outputs.call_args_list
            )
        )

    def test_workbook_writer_uses_final_execution_directory(self) -> None:
        settings = 交易策略回测.build_run_settings(make_case())
        benchmark = 交易策略回测.BenchmarkData(
            code="BENCHMARK",
            name="测试基准",
            closes={},
        )
        with TemporaryDirectory() as temporary_directory:
            output_directory = Path(temporary_directory) / "close"
            with patch.object(
                交易策略回测,
                "write_annual_metrics_workbook",
                return_value=output_directory / "annual.xlsx",
            ) as annual, patch.object(
                交易策略回测,
                "write_backtest_metrics_workbook",
                return_value=output_directory / "backtest.xlsx",
            ) as backtest, patch.object(
                交易策略回测,
                "write_holdings_workbook",
                return_value=output_directory / "holdings.xlsx",
            ) as holdings, patch.object(
                交易策略回测,
                "write_time_series_workbook",
                return_value=output_directory / "series.xlsx",
            ) as time_series, patch.object(
                交易策略回测,
                "write_account_details_workbook",
                return_value=output_directory / "accounts.xlsx",
            ) as accounts:
                交易策略回测.write_mode_workbooks(
                    "close",
                    [],
                    [],
                    [],
                    date(2021, 1, 1),
                    benchmark,
                    [],
                    [],
                    [],
                    "return_r2",
                    output_directory,
                    settings,
                )

            self.assertEqual(annual.call_args.args[3], output_directory)
            self.assertEqual(backtest.call_args.args[4], output_directory)
            self.assertEqual(backtest.call_args.args[6], settings)
            self.assertEqual(holdings.call_args.args[2], output_directory)
            self.assertEqual(time_series.call_args.args[4], output_directory)
            self.assertEqual(accounts.call_args.args[4], output_directory)
            self.assertFalse((output_directory / "close").exists())

    def test_parameter_workbook_records_actual_experiment_case(self) -> None:
        experiment_case = make_case(
            update_frequency="monthly",
            filter_profile="bias_and_volatility",
            correlation_threshold=0.8,
            trend_factor="return_vol",
            trend_window=40,
            stagger_days=3,
        )
        settings = 交易策略回测.build_run_settings(experiment_case)
        benchmark = 交易策略回测.BenchmarkData(
            code="BENCHMARK",
            name="测试基准",
            closes={},
        )
        record = SimpleNamespace(
            rebalance_attempted=True,
            rebalance_succeeded=True,
        )
        performance = {
            "回测开始日": "2021-01-04",
            "回测结束日": "2021-01-05",
            "交易日数量": 2,
            "初始净值": 1.0,
            "期末净值": 1.1,
            "累计收益率": 0.1,
            "年化收益率": 0.1,
            "年化波动率": 0.2,
            "夏普比率": 1.0,
            "Sortino比率": 1.0,
            "最大回撤": -0.1,
            "最大回撤开始日": "2021-01-04",
            "最大回撤结束日": "2021-01-05",
            "Calmar比率": 1.0,
            "年化单边换手率（倍）": 1.0,
            "平均每日单边换手率": 0.1,
            "累计交易成本率": 0.01,
            "平均持仓ETF数量": 1.0,
            "调仓次数": 1,
        }
        relative_performance = {
            "基准年化收益率": 0.05,
            "超额年化收益率": 0.05,
            "跟踪误差": 0.1,
            "信息比率": 0.5,
            "超额最大回撤": -0.05,
            "超额最大回撤开始日": "2021-01-04",
            "超额最大回撤结束日": "2021-01-05",
        }

        with TemporaryDirectory() as temporary_directory, patch.object(
            交易策略回测,
            "calculate_performance",
            return_value=performance,
        ), patch.object(
            交易策略回测,
            "calculate_relative_performance",
            return_value=relative_performance,
        ):
            output_directory = Path(temporary_directory)
            output_path = (
                交易策略回测.write_backtest_metrics_workbook(
                    "close",
                    [record],
                    [],
                    benchmark,
                    output_directory,
                    "return_vol",
                    settings,
                )
            )
            workbook = load_workbook(output_path, data_only=True)
            try:
                parameters = {
                    (row[0], row[1]): row[2]
                    for row in workbook["parameters"].iter_rows(
                        min_row=2,
                        values_only=True,
                    )
                }
            finally:
                workbook.close()

        self.assertEqual(
            parameters[("基本信息", "指数池更新频率")],
            "monthly",
        )
        self.assertEqual(
            parameters[("基本信息", "过滤状态")],
            "bias_and_volatility",
        )
        self.assertEqual(parameters[("趋势策略", "聚类相关性阈值")], 0.8)
        self.assertEqual(parameters[("趋势策略", "趋势因子窗口")], 40)
        self.assertEqual(
            parameters[("趋势策略", "排名得分公式")],
            "收益率÷波动率",
        )
        self.assertEqual(parameters[("波动过滤", "是否启用")], "是")
        self.assertEqual(parameters[("BIAS过滤", "BIAS模式")], "upper")
        self.assertEqual(parameters[("账户结构", "账户数量")], 3)
        self.assertEqual(
            parameters[("账户结构", "单账户调仓间隔（交易日）")],
            3,
        )


if __name__ == "__main__":
    unittest.main()
